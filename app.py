from __future__ import annotations

import os
import uuid
from pathlib import Path

from fastapi import FastAPI, File, Form, Request, UploadFile
from fastapi.responses import FileResponse, HTMLResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from pptx import Presentation

from agents import ReviewOrchestrator

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "uploads"
OUTPUT_DIR = BASE_DIR / "outputs"

UPLOAD_DIR.mkdir(exist_ok=True)
OUTPUT_DIR.mkdir(exist_ok=True)

SUPPORTED_LANGUAGES = {
    "ko": "한국어",
    "en": "영어",
    "pl": "폴란드어",
}

app = FastAPI(title="PPT Translator")
app.mount("/static", StaticFiles(directory=BASE_DIR / "static"), name="static")
templates = Jinja2Templates(directory=str(BASE_DIR / "templates"))


@app.get("/", response_class=HTMLResponse)
def index(request: Request) -> HTMLResponse:
    return templates.TemplateResponse(
        "index.html",
        {
            "request": request,
            "languages": SUPPORTED_LANGUAGES,
            "result": None,
        },
    )


@app.post("/translate", response_class=HTMLResponse)
async def translate(
    request: Request,
    ppt_file: UploadFile = File(...),
    target_language: str = Form(...),
) -> HTMLResponse:
    if target_language not in SUPPORTED_LANGUAGES:
        return templates.TemplateResponse(
            "index.html",
            {
                "request": request,
                "languages": SUPPORTED_LANGUAGES,
                "result": {
                    "error": "지원하지 않는 언어입니다.",
                },
            },
            status_code=400,
        )

    file_id = uuid.uuid4().hex
    input_path = UPLOAD_DIR / f"{file_id}_{ppt_file.filename}"
    output_path = OUTPUT_DIR / f"translated_{file_id}.pptx"

    with input_path.open("wb") as buffer:
        buffer.write(await ppt_file.read())

    presentation = Presentation(str(input_path))
    for slide in presentation.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.text = translate_text(run.text, target_language)

    presentation.save(str(output_path))

    reviewer = ReviewOrchestrator(loop_count=5)
    review_log = reviewer.run(
        {
            "filename": ppt_file.filename,
            "target_language": target_language,
            "slide_count": len(presentation.slides),
        }
    )

    return templates.TemplateResponse(
        "index.html",
        {
            "request": request,
            "languages": SUPPORTED_LANGUAGES,
            "result": {
                "download_name": output_path.name,
                "review_log": review_log,
                "target_language": SUPPORTED_LANGUAGES[target_language],
            },
        },
    )


@app.get("/download/{filename}")
def download(filename: str) -> FileResponse:
    file_path = OUTPUT_DIR / filename
    if not file_path.exists():
        raise FileNotFoundError("Translated file not found.")
    return FileResponse(
        path=file_path,
        filename=filename,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
    )


def translate_text(text: str, target_language: str) -> str:
    if not text.strip():
        return text
    return f"[{SUPPORTED_LANGUAGES[target_language]}] {text}"


if os.getenv("DEBUG"):
    import uvicorn

    uvicorn.run(app, host="0.0.0.0", port=8000)
